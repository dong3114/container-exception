"""
Build Executive-Ready, Linear/Stripe-styled Modern HTML/PDF Proposal
"""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def build_executive_html(output_html):
    html = """<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="utf-8">
<title>CERA Executive Proposal</title>
<style>
  @page {
    size: A4;
    margin: 16mm 14mm 18mm 14mm;
  }
  * {
    box-sizing: border-box;
    -webkit-print-color-adjust: exact;
    print-color-adjust: exact;
  }
  body {
    font-family: 'Pretendard', 'Malgun Gothic', -apple-system, BlinkMacSystemFont, sans-serif;
    color: #0f172a;
    line-height: 1.55;
    font-size: 9.5pt;
    margin: 0;
    padding: 0;
    background-color: #ffffff;
  }
  
  /* Hero Banner */
  .hero {
    background: linear-gradient(135deg, #0b0f19 0%, #1e293b 100%);
    color: #ffffff;
    padding: 26px 28px;
    border-radius: 12px;
    margin-bottom: 22px;
    border: 1px solid #334155;
  }
  .badge {
    display: inline-block;
    font-family: 'Consolas', monospace;
    font-size: 8pt;
    font-weight: 700;
    letter-spacing: 0.8px;
    color: #38bdf8;
    background: rgba(56, 189, 248, 0.12);
    border: 1px solid rgba(56, 189, 248, 0.3);
    padding: 3px 8px;
    border-radius: 9999px;
    margin-bottom: 10px;
    text-transform: uppercase;
  }
  .hero-title {
    font-size: 22pt;
    font-weight: 800;
    letter-spacing: -0.5px;
    margin: 0 0 6px 0;
    line-height: 1.25;
    color: #ffffff;
  }
  .hero-subtitle {
    font-size: 11pt;
    color: #94a3b8;
    margin: 0 0 12px 0;
    font-weight: 400;
  }
  .hero-desc {
    font-size: 9pt;
    color: #cbd5e1;
    line-height: 1.5;
    margin: 0;
    border-top: 1px solid rgba(255, 255, 255, 0.1);
    padding-top: 10px;
  }

  /* Section Styles */
  .section-title {
    font-size: 12.5pt;
    font-weight: 700;
    color: #1e3a8a;
    display: flex;
    align-items: center;
    gap: 8px;
    margin: 20px 0 10px 0;
    border-bottom: 1.5px solid #e2e8f0;
    padding-bottom: 5px;
  }
  .section-title::before {
    content: '';
    display: inline-block;
    width: 4px;
    height: 14px;
    background-color: #2563eb;
    border-radius: 2px;
  }

  /* Big Metric Callout Grid */
  .metric-grid {
    display: grid;
    grid-template-columns: repeat(3, 1fr);
    gap: 12px;
    margin-bottom: 18px;
  }
  .metric-card {
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 8px;
    padding: 14px;
    border-top: 3.5px solid #2563eb;
  }
  .metric-card.danger { border-top-color: #ef4444; }
  .metric-card.success { border-top-color: #10b981; }
  .metric-label {
    font-size: 8pt;
    font-weight: 700;
    color: #64748b;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    margin-bottom: 4px;
  }
  .metric-value {
    font-size: 20pt;
    font-weight: 800;
    letter-spacing: -0.5px;
    color: #0f172a;
    margin-bottom: 4px;
    line-height: 1.1;
  }
  .metric-card.danger .metric-value { color: #dc2626; }
  .metric-card.success .metric-value { color: #059669; }
  .metric-sub {
    font-size: 7.5pt;
    color: #64748b;
    line-height: 1.35;
  }

  /* 2-Column Comparison */
  .compare-grid {
    display: grid;
    grid-template-columns: 1fr 1fr;
    gap: 14px;
    margin-bottom: 18px;
  }
  .compare-card {
    border-radius: 8px;
    padding: 14px;
    font-size: 8.5pt;
  }
  .compare-card.asis {
    background: #fff1f2;
    border: 1px solid #fecdd3;
  }
  .compare-card.tobe {
    background: #f0fdf4;
    border: 1px solid #bbf7d0;
  }
  .compare-header {
    font-size: 10.5pt;
    font-weight: 700;
    margin-bottom: 8px;
    display: flex;
    justify-content: space-between;
  }
  .compare-card.asis .compare-header { color: #be123c; }
  .compare-card.tobe .compare-header { color: #15803d; }
  .compare-list {
    margin: 0;
    padding-left: 16px;
  }
  .compare-list li {
    margin-bottom: 5px;
    color: #334155;
  }

  /* Feature 4-Column Grid */
  .feature-grid {
    display: grid;
    grid-template-columns: repeat(4, 1fr);
    gap: 10px;
    margin-bottom: 18px;
  }
  .feature-card {
    background: #ffffff;
    border: 1px solid #e2e8f0;
    border-radius: 8px;
    padding: 10px 12px;
  }
  .feature-title {
    font-size: 9.5pt;
    font-weight: 700;
    color: #1e293b;
    margin-bottom: 4px;
  }
  .feature-tech {
    font-family: 'Consolas', monospace;
    font-size: 7.5pt;
    font-weight: 700;
    color: #2563eb;
    margin-bottom: 6px;
  }
  .feature-desc {
    font-size: 8pt;
    color: #64748b;
    line-height: 1.4;
  }

  /* Table Style */
  table {
    width: 100%;
    border-collapse: collapse;
    margin: 10px 0;
    font-size: 8pt;
  }
  th {
    background: #f1f5f9;
    color: #0f172a;
    font-weight: 700;
    border: 1px solid #cbd5e1;
    padding: 6px 8px;
    text-align: left;
  }
  td {
    border: 1px solid #e2e8f0;
    padding: 5px 8px;
    color: #334155;
  }
  tr:nth-child(even) td {
    background: #f8fafc;
  }

  /* Footer Meta */
  .footer-meta {
    margin-top: 24px;
    padding-top: 10px;
    border-top: 1px solid #e2e8f0;
    font-size: 7.5pt;
    color: #94a3b8;
    display: flex;
    justify-content: space-between;
  }
</style>
</head>
<body>

  <!-- Hero Header -->
  <div class="hero">
    <div class="badge">Operations Decision Layer</div>
    <div class="hero-title">CERA : Container Exception Recovery Agent</div>
    <div class="hero-subtitle">수입 컨테이너 운영 예외 복구 지능화 및 총비용 최소화 의사결정 플랫폼</div>
    <div class="hero-desc">
      공공 통관 API, 선사·터미널 포털, 계약서, 현장 메신저 등 파편화된 정보를 <strong>신뢰도 기반 단일 운영 상태(Canonical State)</strong>로 통합하고, 창고 입고 거절 시 트럭·ODCY 제약 안에서 <strong>총비용 최소화 복구 계획을 1초 만에 산출</strong>합니다.
    </div>
  </div>

  <!-- Big Number KPI Grid (Squint Test) -->
  <div class="section-title">부산항 30대 Import FCL 실증 성과 (11:00 창고 수용량 급감 예외)</div>
  <div class="metric-grid">
    <div class="metric-card danger">
      <div class="metric-label">방치 시 D&amp;D 노출액</div>
      <div class="metric-value">4,650,000원</div>
      <div class="metric-sub">창고 입고 거절(20→8대)로 12대 적체 시 선사 Free Time 초과 페널티 손실</div>
    </div>
    <div class="metric-card">
      <div class="metric-label">CERA 복구 계획 (Plan v1)</div>
      <div class="metric-value">3,835,000원</div>
      <div class="metric-sub">직송 8대 + ODCY A/B 임시 보관 8대 + 우회 배정 다차원 최적해</div>
    </div>
    <div class="metric-card success">
      <div class="metric-label">순수 손실 방어액 (Saving)</div>
      <div class="metric-value">+815,000원</div>
      <div class="metric-sub">복구 솔버 적용 즉시 방어한 순수 비용 (소요 시간: 3~4시간 → 10분 내 해결)</div>
    </div>
  </div>

  <!-- Paradigm Shift Comparison Grid -->
  <div class="section-title">기존 방식 vs CERA 혁신 패러다임</div>
  <div class="compare-grid">
    <div class="compare-card asis">
      <div class="compare-header">
        <span>기존 단순 현황 조회 (As-Is)</span>
        <span>❌ 한계</span>
      </div>
      <ul class="compare-list">
        <li><strong>단순 위치 추적</strong>: "컨테이너가 어디에 있는가?"만 보여줄 뿐 최적 행동을 제시 못함</li>
        <li><strong>데이터 신뢰도 붕괴</strong>: 공공 API 지연(3~6h) 및 현장 전산 불일치로 사후 수습 반복</li>
        <li><strong>건별 직관 대응</strong>: 한 대 급하다고 ODCY 넣었다가 내일 더 급한 화물의 트럭을 빼앗음</li>
        <li><strong>수기 소통 병목</strong>: 운송사·ODCY 5곳에 일일이 전화·메일 돌리느라 3~4시간 낭비</li>
      </ul>
    </div>
    <div class="compare-card tobe">
      <div class="compare-header">
        <span>CERA 의사결정 플랫폼 (To-Be)</span>
        <span>✓ 혁신</span>
      </div>
      <ul class="compare-list">
        <li><strong>행동 의사결정 레이어</strong>: "지금 무엇을 믿을 수 있고, 무엇을 해야 손실이 최소인가?"</li>
        <li><strong>신뢰도 기반 조율</strong>: 모든 입력을 Event로 다루고 Reconciliation 규칙으로 진실 도출</li>
        <li><strong>수학적 최적화</strong>: 수십 대의 Free Time·보관료·트럭 가용량을 연립하여 최적해 산출</li>
        <li><strong>자동 소통 루프</strong>: 원클릭 문의 초안 → 인간 승인 → 회신 파싱 및 실시간 재최적화</li>
      </ul>
    </div>
  </div>

  <!-- 4-Pillars Engine Boundary -->
  <div class="section-title">시스템 아키텍처 &amp; 신뢰성 설계 (결정론 vs AI 책임 경계)</div>
  <div class="feature-grid">
    <div class="feature-card">
      <div class="feature-title">1. 데이터 조율 엔진</div>
      <div class="feature-tech">Reconciliation Core</div>
      <div class="feature-desc">공공 API, 현장 보고, 메일 출처별 신뢰도 가중치를 부여해 단일 운영 정본(Canonical State) 유지</div>
    </div>
    <div class="feature-card">
      <div class="feature-title">2. 100% 결정론 비용 엔진</div>
      <div class="feature-tech">Tariff &amp; Cost Engine</div>
      <div class="feature-desc">선사별 Free Time 및 D&amp;D 요율 계산은 LLM을 배제하고 순수 수식과 파이썬 코드로 1원 단위 계산</div>
    </div>
    <div class="feature-card">
      <div class="feature-title">3. 다중 복구 최적화 솔버</div>
      <div class="feature-tech">MIP / CP-SAT Solver</div>
      <div class="feature-desc">창고 수용 한계, 게이트 마감, 시간대별 트럭 제약 속에서 글로벌 비용 최소 배정해를 1초 내 도출</div>
    </div>
    <div class="feature-card">
      <div class="feature-title">4. 인간 승인 게이트</div>
      <div class="feature-tech">Human-in-the-Loop</div>
      <div class="feature-desc">대외 메일 발송 및 복구 계획 확정은 운영자의 명시적 승인 없이 절대로 실행되지 않는 안전장치 탑재</div>
    </div>
  </div>

  <!-- Roadmap Summary -->
  <div class="section-title">추진 로드맵 및 단계별 마일스톤</div>
  <table>
    <thead>
      <tr>
        <th style="width: 20%;">단계 (Phase)</th>
        <th style="width: 25%;">목표</th>
        <th>주요 산출물 및 검증 기준</th>
      </tr>
    </thead>
    <tbody>
      <tr>
        <td><strong>Phase 0.5 (완료)</strong></td>
        <td>Control Tower 프로토타입</td>
        <td>Next.js 관제탑 UI 구축 완료, 부산항 30대 시나리오 실증, 개념 데이터 모델(docs/21) 설계 완료</td>
      </tr>
      <tr>
        <td><strong>Phase 1.0 (다음 단계)</strong></td>
        <td>Tech Validation</td>
        <td>Gmail OAuth 연동, 승인된 테스트 계정 실제 발송/회신 파싱 루프 기술 검증, DB 스키마 확정</td>
      </tr>
      <tr>
        <td><strong>Phase 2.0 (실증)</strong></td>
        <td>Pilot 협력사 운영</td>
        <td>중견 포워더 1~2개사 실화물 B/L 50건 대상 관제, 현장 D&amp;D 페널티 방어 실증, Tariff 파서 고도화</td>
      </tr>
      <tr>
        <td><strong>Phase 3.0 (상용화)</strong></td>
        <td>Scale &amp; Commercial</td>
        <td>포워더/화주용 멀티테넌시 SaaS 출시, TMS 연동 OpenAPI 배포, 인천·광양 등 전국 항만 확장</td>
      </tr>
    </tbody>
  </table>

  <div class="footer-meta">
    <span>CERA Project Executive Summary | Confidential</span>
    <span>Powered by CERA Engine Core v0.2</span>
  </div>

</body>
</html>"""

    with open(output_html, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"Generated modern HTML at {output_html}")

def build_executive_docx(output_docx):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Malgun Gothic'
    style_normal.font.size = Pt(9.5)
    style_normal.font.color.rgb = RGBColor(31, 41, 55)

    # Title
    p_badge = doc.add_paragraph()
    r_badge = p_badge.add_run("OPERATIONS DECISION LAYER")
    r_badge.font.name = 'Consolas'
    r_badge.font.size = Pt(9)
    r_badge.font.bold = True
    r_badge.font.color.rgb = RGBColor(37, 99, 235)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("CERA : Container Exception Recovery Agent")
    r_title.font.name = 'Malgun Gothic'
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("수입 컨테이너 운영 예외 복구 지능화 및 총비용 최소화 의사결정 플랫폼")
    r_sub.font.name = 'Malgun Gothic'
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Big Metrics Table (3 Cards)
    t_metrics = doc.add_table(rows=2, cols=3)
    t_metrics.alignment = WD_TABLE_ALIGNMENT.CENTER
    m_headers = ["방치 시 D&D 노출액", "CERA 복구 계획 (Plan v1)", "순수 손실 방어액 (Saving)"]
    m_values = ["4,650,000원", "3,835,000원", "+815,000원"]
    m_colors = [RGBColor(220, 38, 38), RGBColor(15, 23, 42), RGBColor(5, 150, 105)]
    
    for c_idx in range(3):
        # Header cell
        cell_h = t_metrics.rows[0].cells[c_idx]
        p = cell_h.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(m_headers[c_idx])
        r.font.name = 'Malgun Gothic'; r.font.size = Pt(8.5); r.font.bold = True
        r.font.color.rgb = RGBColor(100, 116, 139)
        
        # Value cell
        cell_v = t_metrics.rows[1].cells[c_idx]
        p_v = cell_v.paragraphs[0]
        p_v.paragraph_format.space_before = Pt(2)
        p_v.paragraph_format.space_after = Pt(6)
        r_v = p_v.add_run(m_values[c_idx])
        r_v.font.name = 'Malgun Gothic'; r_v.font.size = Pt(16); r_v.font.bold = True
        r_v.font.color.rgb = m_colors[c_idx]

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 1: Paradigm Shift
    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_before = Pt(12)
    p_s1.paragraph_format.space_after = Pt(6)
    r_s1 = p_s1.add_run("■ 기존 방식 vs CERA 혁신 패러다임 비교")
    r_s1.font.name = 'Malgun Gothic'; r_s1.font.size = Pt(12); r_s1.font.bold = True
    r_s1.font.color.rgb = RGBColor(30, 58, 138)

    t_comp = doc.add_table(rows=5, cols=2)
    t_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_comp.rows[0].cells[0].paragraphs[0].add_run("기존 단순 현황 조회 (As-Is)").font.bold = True
    t_comp.rows[0].cells[1].paragraphs[0].add_run("CERA 의사결정 플랫폼 (To-Be)").font.bold = True
    
    comparisons = [
        ("• 단순 위치 추적: '어디에 있는가?'만 보여줌", "• 행동 의사결정: '무엇을 해야 손실이 최소인가?'"),
        ("• 공공 API 지연(3~6h)으로 사후 수습 반복", "• 신뢰도 기반 조율: 원천 사실과 상태의 엄격한 분리"),
        ("• 직관에 의존한 건별 땜질식 배차", "• 수학적 최적화: 수십 대의 제약을 연립한 최적해 산출"),
        ("• 수기 전화·메일 취합으로 3~4시간 소요", "• 원클릭 초안 → 승인 → 회신 파싱 및 재최적화 (10분)")
    ]
    for r_idx, (asis, tobe) in enumerate(comparisons, start=1):
        t_comp.rows[r_idx].cells[0].paragraphs[0].add_run(asis)
        t_comp.rows[r_idx].cells[1].paragraphs[0].add_run(tobe)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 2: 4-Pillars Engine Boundary
    p_s2 = doc.add_paragraph()
    p_s2.paragraph_format.space_before = Pt(12)
    p_s2.paragraph_format.space_after = Pt(6)
    r_s2 = p_s2.add_run("■ 시스템 아키텍처 & 신뢰성 설계 (결정론 vs AI 책임 경계)")
    r_s2.font.name = 'Malgun Gothic'; r_s2.font.size = Pt(12); r_s2.font.bold = True
    r_s2.font.color.rgb = RGBColor(30, 58, 138)

    pillars = [
        ("1. 데이터 조율 엔진 (Reconciliation Core)", "공공 API, 현장 보고, 메일 등 출처별 신뢰도 가중치를 부여해 단일 운영 정본(Canonical State) 유지."),
        ("2. 100% 결정론 비용 엔진 (Tariff & Cost Engine)", "선사별 Free Time 및 D&D 요율 계산은 LLM을 배제하고 순수 수식과 파이썬 코드로 1원 단위 계산."),
        ("3. 다중 복구 최적화 솔버 (MIP / CP-SAT Solver)", "창고 수용 한계, 게이트 마감, 시간대별 트럭 제약 속에서 글로벌 비용 최소 배정해를 1초 내 도출."),
        ("4. 인간 승인 게이트 (Human-in-the-Loop Gate)", "대외 메일 발송 및 복구 계획 확정은 운영자의 명시적 승인 없이 절대로 실행되지 않는 안전장치 탑재.")
    ]
    for p_title, p_desc in pillars:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(3)
        r_t = p.add_run(p_title + "\n"); r_t.font.bold = True
        r_d = p.add_run(p_desc)
        r_d.font.color.rgb = RGBColor(71, 85, 105)

    doc.save(output_docx)
    print(f"Generated modern DOCX at {output_docx}")

if __name__ == '__main__':
    out_html = r"C:\Users\PC\Desktop\workspace\container_exception\data\Proposal\01_EXECUTIVE_PROPOSAL.html"
    out_docx = r"C:\Users\PC\Desktop\workspace\container_exception\data\Proposal\01_EXECUTIVE_PROPOSAL.docx"
    build_executive_html(out_html)
    build_executive_docx(out_docx)
